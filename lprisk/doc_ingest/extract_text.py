"""IM 텍스트·표 추출 (설계서 §3.5 extract_pdf_text / extract_pptx, DOCX 추가).

출력은 im-analyst가 읽는 단위(페이지·슬라이드·문단·표) 목록이다. 텍스트 레이어가 없는 페이지는
PNG로 렌더링해 경로를 남긴다(im-analyst가 비전으로 읽음).
"""
from __future__ import annotations

from pathlib import Path

import fitz

from ..io import now_iso, save_json
from ..paths import ROOT, rel, wf2_dir
from .detect_format import SCAN_CHARS_PER_PAGE

RENDER_DPI = 110


def _table_rows(tab) -> list[list[str]]:
    rows = tab.extract()
    return [[(c or "").strip() for c in r] for r in rows]


def pdf_units(path: Path, render_dir: Path | None) -> list[dict]:
    doc = fitz.open(path)
    units = []
    for i, page in enumerate(doc):
        text = page.get_text().strip()
        unit = {"locator": {"page": i + 1}, "text": text, "tables": []}
        try:
            for tab in page.find_tables().tables:
                rows = _table_rows(tab)
                if rows and any(any(c for c in r) for r in rows):
                    unit["tables"].append(rows)
        except Exception:
            pass
        if len(text) < SCAN_CHARS_PER_PAGE and render_dir is not None:
            render_dir.mkdir(parents=True, exist_ok=True)
            out = render_dir / f"{path.stem[:40]}_p{i + 1:03d}.png"
            if not out.exists():
                page.get_pixmap(dpi=RENDER_DPI).save(out)
            unit["image"] = rel(out)
            unit["scan"] = True
        units.append(unit)
    return units


def pptx_units(path: Path) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(path)
    units = []
    for i, slide in enumerate(prs.slides):
        texts, tables = [], []
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                texts.append(shp.text_frame.text.strip())
            if getattr(shp, "has_table", False) and shp.has_table:
                tables.append([[c.text.strip() for c in r.cells] for r in shp.table.rows])
        units.append({"locator": {"slide": i + 1}, "text": "\n".join(texts), "tables": tables})
    return units


def _docx_rows(tbl) -> tuple[list[list[str]], list[tuple[int, object]]]:
    """표 행 텍스트와 중첩 표 목록 [(행 번호, 중첩 표)]. 병합 셀은 한 번만."""
    rows, nested = [], []
    for ri, r in enumerate(tbl.rows):
        cells, prev = [], None
        for c in r.cells:  # 병합 셀은 같은 객체가 반복됨
            if c._tc is prev:
                continue
            prev = c._tc
            cells.append(c.text.strip())
            nested += [(ri, nt) for nt in c.tables]
        rows.append(cells)
    return rows, nested


def _docx_images(doc, element, out_dir: Path | None, stem: str, tag: str) -> list[str]:
    """요소 안의 그림(a:blip)을 PNG/원본 형식으로 저장. 반환: 상대 경로 목록."""
    if out_dir is None:
        return []
    paths = []
    for i, blip in enumerate(element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")):
        rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        part = doc.part.related_parts.get(rid) if rid else None
        if part is None:
            continue
        ext = Path(part.partname).suffix or ".png"
        out_dir.mkdir(parents=True, exist_ok=True)
        fn = out_dir / f"{stem[:30]}_{tag}_img{i + 1}{ext}"
        if not fn.exists():
            fn.write_bytes(part.blob)
        paths.append(rel(fn))
    return paths


def docx_units(path: Path, image_dir: Path | None = None) -> list[dict]:
    """문단은 묶음 번호로, 표는 최상위 표 번호(1부터)로 기록한다.

    표 안의 중첩 표는 같은 단위의 `tables`에 이어 붙이고 `nested`에 (행, 순번)을 남긴다.
    문단·표 안의 그림은 파일로 꺼내 `images`에 경로를 남긴다(im-analyst가 비전으로 읽음).
    """
    import docx
    d = docx.Document(path)
    units = []
    body = d.element.body
    p_idx = t_idx = 0
    buf, buf_start, buf_imgs = [], None, []

    def flush():
        nonlocal buf, buf_start, buf_imgs
        if buf or buf_imgs:
            u = {"locator": {"para": buf_start, "para_end": p_idx}, "text": "\n".join(buf), "tables": []}
            if buf_imgs:
                u["images"] = buf_imgs
            units.append(u)
        buf, buf_start, buf_imgs = [], None, []

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p_idx += 1
            text = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t")).strip()
            imgs = _docx_images(d, child, image_dir, path.stem, f"p{p_idx}")
            if (text or imgs) and buf_start is None:
                buf_start = p_idx
            if text:
                buf.append(f"[¶{p_idx}] {text}")
            buf_imgs += imgs
            if len(buf) >= 25:
                flush()
        elif tag == "tbl":
            flush()
            t_idx += 1
            tbl = d.tables[t_idx - 1]
            rows, nested = _docx_rows(tbl)
            tables, nest_meta = [rows], []
            for k, (ri, nt) in enumerate(nested):
                nrows, _ = _docx_rows(nt)
                tables.append(nrows)
                nest_meta.append({"row": ri, "nested": k + 1})
            u = {"locator": {"table": t_idx, "after_para": p_idx}, "text": "", "tables": tables}
            if nest_meta:
                u["nested"] = nest_meta
            imgs = _docx_images(d, child, image_dir, path.stem, f"t{t_idx}")
            if imgs:
                u["images"] = imgs
            units.append(u)
    flush()
    return units


def run(asset_id: str, inventory: dict) -> dict:
    out_dir = wf2_dir(asset_id)
    docs = []
    for f in inventory["im"]:
        if not f.get("readable"):
            continue
        path = ROOT / f["path"]
        fmt = f["format"]
        if fmt.startswith("pdf"):
            units = pdf_units(path, out_dir / "pages")
        elif fmt == "pptx":
            units = pptx_units(path)
        elif fmt == "docx":
            units = docx_units(path, out_dir / "pages")
        else:
            continue
        docs.append({"path": f["path"], "format": fmt, "units": units})
    data = {"asset_id": asset_id, "created_at": now_iso(), "docs": docs}
    save_json(out_dir / "w2_02_im_text.json", data)
    return data
